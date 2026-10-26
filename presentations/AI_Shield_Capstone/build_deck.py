"""Build matching editable PowerPoint and PDF slides from local project evidence."""
from pathlib import Path
import json
import math
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


ROOT = Path(__file__).resolve().parent
DATA = json.loads((ROOT / "slide_content.json").read_text(encoding="utf-8"))
W, H = 960, 540
INK, MUTED, GREEN, RED, GOLD = "172C2B", "546966", "00796B", "C84760", "B77B00"
PAPER, LIGHT, LINE, MINT = "FFFFFF", "F2F6F5", "DCE5E2", "DFF2EA"
FONT_DIR = Path("C:/Windows/Fonts")
pdfmetrics.registerFont(TTFont("Arial", str(FONT_DIR / "arial.ttf")))
pdfmetrics.registerFont(TTFont("ArialBold", str(FONT_DIR / "arialbd.ttf")))


def rgb(value):
    return tuple(int(value[i:i + 2], 16) / 255 for i in (0, 2, 4))


def wrap(text, width, size, bold=False):
    font = "ArialBold" if bold else "Arial"
    lines = []
    for paragraph in str(text).split("\n"):
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}".strip()
            if pdfmetrics.stringWidth(candidate, font, size) <= width:
                current = candidate
            else:
                if not current:
                    raise ValueError(f"Unbreakable word exceeds text box: {word}")
                lines.append(current)
                current = word
        lines.append(current)
    return lines


class Deck:
    def __init__(self):
        self.ppt = Presentation()
        self.ppt.slide_width, self.ppt.slide_height = Pt(W), Pt(H)
        self.ppt.core_properties.title = "AI Shield: Capstone Research and Development"
        self.ppt.core_properties.author = DATA["author"]
        self.ppt.core_properties.subject = "A milestone-based project history without calendar dates"
        self.pdf = canvas.Canvas(str(ROOT / "AI_Shield_Capstone.pdf"), pagesize=(W, H))
        self.pdf.setTitle(self.ppt.core_properties.title)
        self.pdf.setAuthor(DATA["author"])
        self.index = 0
        self.audit = []

    def rect(self, x, y, w, h, fill=LIGHT, stroke=None, radius=0):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Pt(x), Pt(y), Pt(w), Pt(h),
        )
        if radius:
            shape.adjustments[0] = min(0.12, radius / max(1, min(w, h)))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if stroke:
            shape.line.color.rgb = RGBColor.from_string(stroke)
            shape.line.width = Pt(0.75)
        else:
            shape.line.fill.background()
        self.pdf.setFillColorRGB(*rgb(fill))
        self.pdf.setStrokeColorRGB(*rgb(stroke or fill))
        if radius:
            self.pdf.roundRect(x, H - y - h, w, h, radius, fill=1, stroke=bool(stroke))
        else:
            self.pdf.rect(x, H - y - h, w, h, fill=1, stroke=bool(stroke))

    def line(self, x1, y1, x2, y2, color=LINE, width=1):
        from pptx.enum.shapes import MSO_CONNECTOR
        shape = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2))
        shape.line.color.rgb = RGBColor.from_string(color)
        shape.line.width = Pt(width)
        self.pdf.setStrokeColorRGB(*rgb(color))
        self.pdf.setLineWidth(width)
        self.pdf.line(x1, H - y1, x2, H - y2)

    def arrow(self, x1, y1, x2, y2, color=GREEN):
        self.line(x1, y1, x2, y2, color, 1.6)
        angle = math.atan2(y2 - y1, x2 - x1)
        for offset in (-0.5, 0.5):
            self.line(x2, y2, x2 - 7 * math.cos(angle + offset), y2 - 7 * math.sin(angle + offset), color, 1.6)

    def text(self, text, x, y, w, size=18, color=INK, bold=False, h=None, url=None, align="left"):
        lines = wrap(text, w - 2, size, bold)
        leading = size * 1.22
        needed = leading * len(lines) + 3
        if h is not None and needed > h + 0.1:
            raise ValueError(f"Slide {self.index}: text exceeds height ({needed:.1f}>{h}): {text[:75]}")
        if x < 0 or x + w > W + 0.1 or y < 0 or y + needed > H:
            raise ValueError(f"Slide {self.index}: text outside slide: {text[:70]}")
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(needed))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        from pptx.enum.text import PP_ALIGN
        for index, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = Pt(leading)
            paragraph.alignment = PP_ALIGN.CENTER if align == "center" else PP_ALIGN.LEFT
            run = paragraph.add_run()
            run.text = line
            run.font.name, run.font.size, run.font.bold = "Arial", Pt(size), bold
            run.font.color.rgb = RGBColor.from_string(color)
            if url:
                run.hyperlink.address = url
        font = "ArialBold" if bold else "Arial"
        self.pdf.setFont(font, size)
        self.pdf.setFillColorRGB(*rgb(color))
        ascent = pdfmetrics.getAscent(font) * size / 1000
        for index, line in enumerate(lines):
            xx = x + (w - pdfmetrics.stringWidth(line, font, size)) / 2 if align == "center" else x
            self.pdf.drawString(xx, H - y - ascent - index * leading - 1, line)
        if url:
            self.pdf.linkURL(url, (x, H - y - needed, x + w, H - y), relative=0)
        self.audit.append({"slide": self.index, "text": text, "box": [x, y, w, needed], "size": size})
        return needed

    def new(self, stage, title, subtitle="", sources=(0,)):
        if self.index:
            self.pdf.showPage()
        self.index += 1
        self.slide = self.ppt.slides.add_slide(self.ppt.slide_layouts[6])
        self.rect(0, 0, W, H, PAPER)
        self.rect(0, 0, 8, H, GREEN)
        self.text(stage.upper(), 44, 24, 800, 10, GREEN, True)
        self.text(f"{self.index:02d} / 15", 862, 24, 58, 10, MUTED)
        self.text(title, 44, 60, 872, 30, bold=True, h=82)
        if subtitle:
            self.text(subtitle, 46, 111, 870, 14, MUTED, h=38)
        self.line(44, 502, 916, 502)
        self.text("AI SHIELD  /  ABDELRAHMAN SOLIMAN", 44, 515, 400, 8.5, MUTED)
        x = 565
        for source in sources:
            name, url = DATA["sources"][source]
            label = {0: "Project notes", 1: "Build history", 2: "Results", 3: "Dataset notebook", 4: "Training notebook", 5: "Comparison", 6: "Fusion study", 7: "Freeze checks", 8: "Saved artifacts", 9: "Repository"}[source]
            self.text(label, x, 515, 115, 8.5, MUTED, url=url)
            x += 116
        item = DATA["slides"][self.index - 1]
        notes = item["notes"] + "\n\nSources:\n" + "\n".join(f"{DATA['sources'][i][0]}: {DATA['sources'][i][1]}" for i in sources)
        self.slide.notes_slide.notes_text_frame.text = notes

    def bullet(self, title, body, x, y, w, color=GREEN, size=17):
        self.rect(x, y + 4, 4, 18, color)
        height = self.text(title, x + 16, y, w - 16, size, bold=True)
        self.text(body, x + 16, y + height + 7, w - 16, size - 1, MUTED)

    def band(self, text, y=440, color=GREEN, fill=MINT, size=17):
        self.rect(44, y, 872, 45, fill)
        self.text(text, 60, y + 11, 840, size, color, True, h=29)

    def save(self):
        self.ppt.save(ROOT / "AI_Shield_Capstone.pptx")
        self.pdf.save()
        (ROOT / "layout_audit.json").write_text(json.dumps(self.audit, indent=2), encoding="utf-8")
        notes = []
        for index, slide in enumerate(DATA["slides"], 1):
            notes.append(f"SLIDE {index:02d}: {slide['title']}\n\n{slide['notes']}\n")
        notes.append("SOURCE REGISTER\n" + "\n".join(f"{name}: {url}" for name, url in DATA["sources"]))
        (ROOT / "speaker_notes.txt").write_text("\n\n".join(notes), encoding="utf-8")


def build():
    d = Deck()
    # 1: Opening, with a diagram of the actual current concept.
    d.new("Capstone / research and engineering", "AI Shield", sources=(0,))
    d.text("From image detection to\nevidence-aware verification", 46, 141, 470, 33, bold=True, h=100)
    d.text("Abdelrahman Soliman", 48, 270, 440, 20, GREEN, True)
    d.text("A project story: what I built, what failed,\nwhat improved, and where the work stands.", 48, 311, 440, 18, MUTED)
    d.rect(550, 143, 366, 276, LIGHT)
    d.text("ONE IMAGE", 574, 165, 300, 13, MUTED, True)
    d.rect(575, 205, 315, 64, MINT)
    d.text("AI-generation detection", 595, 224, 275, 20, GREEN, True)
    d.rect(575, 285, 315, 87, PAPER, LINE)
    d.text("Recorded origin and\nediting evidence", 595, 302, 275, 20, bold=True)
    d.text("Research evidence, followed by a working local service.", 48, 452, 850, 17, MUTED)

    # 2: Define the problem and avoid conflating detection with provenance.
    d.new("Research question", "Can the prediction survive a change in source?", "The goal expanded from a classifier to an evidence-aware verification tool.")
    for x, n, title, body in [
        (44, "01", "Detect", "Estimate whether an image shows visual evidence of AI generation."),
        (344, "02", "Generalize", "Test on new generators and image sources, not only familiar examples."),
        (644, "03", "Explain evidence", "Show available camera, editing and signed-history information."),
    ]:
        d.text(n, x, 179, 270, 38, GREEN, True)
        d.line(x, 237, x + 266, 237)
        d.text(title, x, 258, 266, 23, bold=True)
        d.text(body, x, 306, 266, 19, MUTED, h=115)
    d.band("An AI score is a prediction. Recorded history is separate evidence.")

    # 3: Date-free development timeline.
    d.new("Project journey", "Six stages, each prompted by a limitation", "Milestones show the development order without calendar dates.", sources=(1, 0))
    stages = [
        ("Face detector", "A working prototype, but a narrow image domain."),
        ("Forensic branches", "Broader visual signals, reconstruction and face routing."),
        ("Semantic + fusion", "Custom patch attention and a learned score combiner."),
        ("Generalization audit", "A newer benchmark exposed the original model's weakness."),
        ("Community Forensics", "A stronger detector, selected by controlled comparison."),
        ("Provenance + service", "Repeatable inference and a combined local upload API."),
    ]
    d.line(63, 174, 63, 452, GREEN, 2)
    for i, (title, body) in enumerate(stages):
        y = 164 + i * 49
        d.rect(47, y, 32, 32, GREEN)
        d.text(str(i + 1), 48, y + 6, 30, 16, PAPER, True, align="center")
        d.text(title, 100, y + 3, 260, 18, bold=True)
        d.text(body, 374, y + 4, 534, 16, MUTED)

    # 4: Initial face prototype and its measured generalization drop.
    d.new("Stage 1 / first prototype", "The first detector worked, but its scope was narrow", "EfficientNetV2-B0 trained on OpenForensics face-manipulation images.", sources=(1, 2))
    d.bullet("What I built", "Training, evaluation and an image-upload demo for face manipulation.", 44, 173, 340)
    d.bullet("What I learned", "A strong validation result did not guarantee the same reliability on held-out images.", 44, 287, 340, RED)
    chart_x, chart_y, chart_h = 450, 176, 232
    d.line(chart_x, chart_y + chart_h, 914, chart_y + chart_h)
    for x, value, label, color in [(502, 96.51, "Validation", GREEN), (716, 85.77, "Held-out test", RED)]:
        height = chart_h * value / 100
        d.rect(x, chart_y + chart_h - height, 127, height, color)
        d.text(f"{value:.2f}%", x - 18, chart_y + chart_h - height - 35, 164, 23, color, True, align="center")
        d.text(label, x - 17, 418, 166, 15, MUTED, align="center")
    d.text("Accuracy; chart baseline = 0%", 450, 457, 450, 12, MUTED)

    # 5: Broader low-level forensic signals, including a negative result.
    d.new("Stage 2 / broader visual evidence", "Built several forensic signals, then tested their value", "Forensic = patterns in the pixels and image-processing traces.", sources=(1, 2))
    branches = [
        (44, "EfficientNet-B0", "General image classifier", "Trained on balanced Tiny-GenImage; calibrated scores and saved branch outputs.", GREEN),
        (344, "Real-only VAE", "Reconstruction test", "Built the encoder, decoder and training loop; tested whether unusual reconstruction errors flag AI.", RED),
        (644, "FaceGate", "Use the face model only when relevant", "Checked face presence and routed the OpenForensics score only when a valid face was found.", GOLD),
    ]
    for x, title, sub, body, color in branches:
        d.rect(x, 173, 266, 5, color)
        d.text(title, x, 200, 266, 23, color, True)
        d.text(sub, x, 245, 266, 16, bold=True, h=60)
        d.text(body, x, 307, 266, 17, MUTED, h=120)
    d.band("Negative finding: VAE alone AUROC ~0.470; a naive weighted blend hurt performance.", size=15)

    # 6: Highlight the student's custom ML component, not just pretrained models.
    d.new("Stage 3 / semantic layer", "Built a custom attention head over image patches", "ViT = Vision Transformer; it represents an image as a sequence of patches.", sources=(1, 2))
    nodes = [(44, 130, "Image\npatches"), (216, 166, "Pretrained\nViT backbone"), (424, 285, "Custom attention\n+ refinement + pooling"), (751, 165, "Whole-image\nAI score")]
    for x, width, label in nodes:
        d.rect(x, 189, width, 98, MINT if x == 424 else LIGHT)
        d.text(label, x + 12, 213, width - 24, 18, GREEN if x == 424 else INK, True, align="center", h=59)
    for left, right in zip(nodes, nodes[1:]):
        d.arrow(left[0] + left[1] + 7, 237, right[0] - 9, 237)
    d.text("856,578", 48, 338, 330, 37, GREEN, True)
    d.text("trainable custom-head parameters", 50, 387, 390, 17, MUTED)
    d.text("Backbone frozen; custom head trained.\nImplemented multi-head attention,\nrefinement and attention pooling.", 470, 339, 440, 18, MUTED, h=100)
    d.band("Visual modeling, not demonstrated caption checking or human-like reasoning.", size=16)

    # 7: Historical fusion benchmark. Explicitly preserve its evaluation universe.
    d.new("Stage 3 / original combined system", "Learned fusion improved the original benchmark", "Original Tiny-GenImage fusion test: 1,400 images. These are historical results.", sources=(1, 2))
    metrics = [("Calibrated EfficientNet", .9558, GOLD), ("Semantic ViT", .9303, RED), ("Forensic-only GLM", .9578, MUTED), ("Full GLM fusion", .9819, GREEN)]
    for i, (name, value, color) in enumerate(metrics):
        y = 170 + i * 57
        d.text(name, 44, y, 250, 15, bold=i == 3)
        d.rect(306, y + 3, 283, 15, LIGHT)
        d.rect(306, y + 3, 283 * value, 15, color)
        d.text(f"{value:.4f}", 603, y - 3, 85, 17, color, True)
    d.text("17 features", 726, 174, 185, 28, GREEN, True)
    d.text("11 forensic + 6 semantic", 727, 218, 185, 16, MUTED)
    d.text("94.07%", 726, 276, 185, 29, bold=True)
    d.text("accuracy on this test", 727, 318, 185, 15, MUTED)
    d.text("GLM = logistic regression: a learned score combiner.\nAUROC measures ranking: 0.5 is chance; 1.0 is perfect. It is not accuracy.", 46, 407, 864, 14, MUTED)
    d.band("Caught 54 additional generated images with only 1 additional false positive.", y=456, size=16)

    # 8: Generalization audit: do not imply scores from distinct datasets are comparable training steps.
    d.new("Stage 4 / turning point", "Newer images exposed a generalization problem", "Re-ran the saved original AI Shield pipeline, without retraining it.", sources=(2, 5))
    d.rect(44, 167, 412, 218, LIGHT)
    d.rect(504, 167, 412, 218, "FBEFF2")
    d.text("ORIGINAL BENCHMARK", 66, 188, 365, 12, MUTED, True)
    d.text("0.9819", 66, 227, 365, 54, GREEN, True)
    d.text("AUROC on the original\nTiny-GenImage fusion test", 68, 312, 361, 18, MUTED)
    d.text("NEW DATASET v1", 526, 188, 365, 12, RED, True)
    d.text("~0.606", 526, 227, 365, 54, RED, True)
    d.text("AUROC on the new ordinary test\n~0.599 on held-out GPT Image 2", 528, 312, 361, 18, MUTED)
    d.text("Different test sets: this contrast shows the transfer problem, not a training curve.", 47, 403, 864, 14, MUTED)
    d.band("Next: compare old AI Shield with Community Forensics (CF) on the same images.", y=450, size=16)

    # 9: Dataset and held-out generator protocol.
    d.new("Stage 4 / controlled benchmark", "Built Dataset v1 to test a harder question", "Fixed image identities and splits; evaluated a generator excluded from our fine-tuning.", sources=(3, 2))
    d.text("5,124", 44, 168, 300, 62, GREEN, True)
    d.text("balanced images", 48, 251, 270, 21, bold=True)
    d.text("2,562 real\n2,562 generated", 48, 297, 270, 20, MUTED)
    segments = [("Training", 2990, GREEN), ("Validation", 640, MUTED), ("Ordinary test", 640, RED), ("Unseen generator", 854, GOLD)]
    x = 358
    for label, count, color in segments:
        width = 558 * count / 5124
        d.rect(x, 178, width - 2, 53, color)
        x += width
    for i, (label, count, color) in enumerate(segments):
        xx, yy = 358 + (i % 2) * 287, 258 + (i // 2) * 72
        d.rect(xx, yy + 4, 10, 10, color)
        d.text(f"{label}: {count:,}", xx + 20, yy, 250, 16, bold=True)
    d.text("Five generator families in development.\nGPT Image 2 reserved for a separate test.", 359, 414, 555, 17, MUTED)
    d.text("Held out from this project's training/validation; no claim about the pretrained model's full history.", 46, 474, 866, 11, MUTED)

    # 10: Validation-led two-stage adaptation.
    d.new("Stage 5 / adapting Community Forensics", "Fine-tuned the detector in two controlled stages", "The checkpoint and decision threshold were selected on validation data.", sources=(4, 2))
    d.bullet("Stage 1: classifier head", "Three epochs; keep the pretrained visual backbone fixed.", 44, 176, 343)
    d.bullet("Stage 2: final blocks", "Three epochs; train the final two transformer blocks, normalization and head.", 44, 294, 343)
    x0, y0, width, height = 450, 184, 437, 205
    for value in (.5, .75, 1.0):
        y = y0 + height * (1 - (value - .5) / .5)
        d.line(x0, y, x0 + width, y)
        d.text(f"{value:.2f}", x0 - 43, y - 8, 37, 11, MUTED)
    points = [(490, .8780, "Baseline"), (665, .9804, "Stage 1"), (840, .9870, "Stage 2")]
    last = None
    for x, value, label in points:
        y = y0 + height * (1 - (value - .5) / .5)
        if last:
            d.line(last[0], last[1], x, y, GREEN, 2.5)
        d.rect(x - 5, y - 5, 10, 10, GREEN)
        d.text(f"{value:.4f}", x - 60, y - 38, 120, 21, GREEN, True, align="center")
        d.text(label, x - 65, 409, 130, 14, MUTED, align="center")
        last = x, y
    d.band("Validation AUROC improved from 0.8780 to 0.9870.", y=451)

    # 11: Same-image comparison of all candidate baselines.
    d.new("Stage 5 / same-image evaluation", "Fine-tuned CF became the strongest detector", "Dataset v1 final evaluation. AUROC is ranking performance, not percentage accuracy.", sources=(2, 4, 5))
    d.text("MODEL", 44, 167, 280, 12, MUTED, True)
    d.text("ORDINARY TEST AUROC", 365, 167, 255, 12, MUTED, True)
    d.text("GPT IMAGE 2 AUROC", 659, 167, 255, 12, MUTED, True)
    rows = [("Old AI Shield fusion", .606, .599, RED), ("Original CF", .8860, .9555, GOLD), ("Fine-tuned CF", .9884, .9946, GREEN)]
    for i, (name, ordinary, unseen, color) in enumerate(rows):
        y = 205 + i * 68
        if i == 2:
            d.rect(44, y - 8, 872, 57, MINT)
        d.text(name, 56, y + 5, 285, 18, color, True)
        for x, value in ((367, ordinary), (662, unseen)):
            d.rect(x, y + 13, 153, 15, LINE)
            d.rect(x, y + 13, 153 * value, 15, color)
            d.text(f"{value:.4f}", x + 166, y + 4, 90, 19, color, True)
    d.text("Fine-tuned accuracy: 95.16% ordinary test | 96.60% held-out-generator test.", 47, 400, 863, 16, MUTED)
    d.rect(44, 433, 872, 61, MINT)
    d.text("Real-image false alarms: 20.63% -> 3.75% on the ordinary test.", 60, 443, 840, 17, GREEN, True)
    d.text("Each CF model uses its validation-selected max-F1 cutoff; bar scales run from 0 to 1.", 60, 477, 840, 9, MUTED)

    # 12: Latest fusion was rejected on evidence.
    d.new("Stage 5 / test the added complexity", "More models did not improve this detector", "New fusion: fine-tuned CF + calibrated EfficientNet + FaceGate/OpenForensics.", sources=(6, 2))
    d.text("MODEL", 58, 183, 410, 12, MUTED, True)
    d.text("TEST AUROC", 504, 183, 160, 12, MUTED, True)
    d.text("GPT IMAGE 2 AUROC", 708, 183, 195, 12, MUTED, True)
    for y, name, first, second, color in [(235, "Fine-tuned CF alone", ".9884", ".9946", GREEN), (318, "New fusion", ".9854", ".9936", MUTED)]:
        d.rect(44, y - 13, 872, 68, MINT if color == GREEN else LIGHT)
        d.text(name, 60, y + 1, 420, 23, color, True)
        d.text("0" + first, 505, y - 1, 160, 28, color, True)
        d.text("0" + second, 709, y - 1, 190, 28, color, True)
    d.text("The combination was slightly worse on both sets.\nThis is a separate experiment from the original 17-feature GLM.", 48, 409, 860, 16, MUTED)
    d.band("Decision: select CF alone; pause further model-fusion experiments.", y=455, size=17)

    # 13: Reproducibility is a concrete engineering contribution.
    d.new("Stage 6 / repeatable inference", "Locked the detector before building around it", "Frozen = same weights, same image preparation, same decision cutoff.", sources=(7, 8))
    checks = [
        ("01", "Checkpoint hash", "A digital fingerprint confirms the exact fine-tuned model file."),
        ("02", "Exact preprocessing", "Orientation -> RGB -> resize -> crop -> normalization."),
        ("03", "Fixed threshold", "About 0.634, selected using validation F1 (precision/recall balance)."),
        ("04", "Reference-image checks", "Six saved examples matched scores within tolerance and kept their labels."),
    ]
    for i, (n, title, body) in enumerate(checks):
        y = 169 + i * 68
        d.rect(44, y, 872, 55, LIGHT)
        d.text(n, 60, y + 12, 47, 20, GREEN, True)
        d.text(title, 119, y + 14, 255, 18, bold=True)
        d.text(body, 385, y + 12, 510, 16, MUTED, h=42)
    d.text("These are implementation-consistency checks, not a new accuracy benchmark.", 48, 464, 860, 14, MUTED)

    # 14: Current implemented service, not an imagined frontend.
    d.new("Stage 6 / current local system", "Added recorded provenance beside the AI score", "FastAPI = the Python service that runs both checks for the future website.", sources=(0, 7))
    d.rect(44, 247, 150, 79, LIGHT)
    d.text("Original\nupload", 57, 263, 124, 20, bold=True, align="center")
    d.rect(281, 171, 328, 86, MINT)
    d.text("CF detector", 299, 186, 290, 21, GREEN, True)
    d.text("Visual AI-generation score", 299, 221, 290, 15, MUTED)
    d.rect(281, 300, 328, 108, LIGHT)
    d.text("ExifTool + C2PA", 299, 315, 290, 21, bold=True)
    d.text("Stored image details + signed\ncreation/editing record checks", 299, 350, 290, 16, MUTED)
    d.rect(709, 245, 207, 94, GREEN)
    d.text("Combined\nevidence report", 725, 264, 175, 21, PAPER, True, align="center")
    for coords in [(194, 274, 275, 215), (194, 302, 275, 353), (609, 215, 703, 274), (609, 353, 703, 310)]:
        d.arrow(*coords)
    d.text("38 tests passed | 6 reference images matched through the live upload API.", 48, 434, 860, 17, GREEN, True)
    d.text("EXIF = stored camera/software/date details. C2PA = signed creation/editing records.", 48, 465, 860, 11, MUTED)
    d.text("Missing metadata = unknown. Photoshop is not proof of AI. Credential validity is not scene truth.", 48, 483, 860, 10, MUTED)

    # 15: Honest current state and a concrete professor discussion.
    d.new("Current state / discussion", "Working locally; evaluation and deployment come next", "The hosted Hugging Face demo still uses the older system.", sources=(0, 9))
    columns = [
        (44, "WORKING NOW", GREEN, "Fine-tuned CF detector\nRepeatable inference\nEXIF / C2PA reporting\nCombined local upload API"),
        (344, "NEXT", GOLD, "JPEG / resize robustness\nArtwork and edited-photo tests\nPublish code + build website\nPrepare public deployment"),
        (644, "LATER", MUTED, "Image / claim consistency\nShort-video support\nBroader source evaluation"),
    ]
    for x, title, color, body in columns:
        d.rect(x, 172, 268, 5, color)
        d.text(title, x, 198, 268, 16, color, True)
        d.text(body, x, 244, 268, 18, MUTED, h=145)
    d.rect(44, 406, 872, 83, MINT)
    d.text("For discussion", 60, 419, 840, 16, GREEN, True)
    d.text("Which evaluation would make the capstone most convincing?\nHow should uncertain or conflicting evidence be presented?", 60, 447, 840, 16, INK, h=42)

    if d.index != 15:
        raise RuntimeError("Expected exactly 15 slides")
    d.save()


def render_and_check():
    import pymupdf
    from PIL import Image, ImageOps, ImageDraw

    document = pymupdf.open(ROOT / "AI_Shield_Capstone.pdf")
    if len(document) != 15:
        raise RuntimeError("PDF slide count differs")
    previews = ROOT / "slide_previews"
    previews.mkdir(exist_ok=True)
    sheet = Image.new("RGB", (1200, 3 * 250), "#dce5e2")
    draw = ImageDraw.Draw(sheet)
    findings = []
    forbidden = re.compile(r"\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\b|\b20\d{2}\b|\bWeek\s+\d+\b")
    for index, page in enumerate(document):
        text = page.get_text()
        if forbidden.search(text):
            raise RuntimeError(f"Calendar date found on slide {index + 1}")
        pix = page.get_pixmap(matrix=pymupdf.Matrix(1.5, 1.5), alpha=False)
        path = previews / f"slide_{index + 1:02d}.png"
        pix.save(path)
        image = Image.open(path).convert("RGB")
        tile = ImageOps.contain(image, (384, 216))
        # Four columns by four rows, leaving a numbered label below each preview.
        if index == 0:
            sheet = Image.new("RGB", (1600, 4 * 250), "#dce5e2")
            draw = ImageDraw.Draw(sheet)
        x, y = (index % 4) * 400 + 8, (index // 4) * 250 + 8
        sheet.paste(tile, (x, y))
        draw.text((x + 4, y + 220), f"{index + 1:02d}  {DATA['slides'][index]['title'][:47]}", fill="#172c2b")
        if len(text.strip()) < 80:
            raise RuntimeError(f"Slide {index + 1} is unexpectedly blank")
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    x0, y0, x1, y1 = span["bbox"]
                    if x0 < -1 or y0 < -1 or x1 > W + 1 or y1 > H + 1:
                        findings.append({"slide": index + 1, "text": span["text"], "issue": "out of bounds"})
    sheet.save(ROOT / "slide_overview.png")
    presentation = Presentation(ROOT / "AI_Shield_Capstone.pptx")
    if len(presentation.slides) != 15 or not all(slide.notes_slide.notes_text_frame.text.strip() for slide in presentation.slides):
        raise RuntimeError("PowerPoint slides or speaker notes are missing")
    if findings:
        raise RuntimeError(str(findings))
    report = {"slides": 15, "powerpoint_editable": True, "speaker_notes": True,
              "calendar_dates_on_slides": False, "out_of_bounds_text": findings,
              "generation": "Local vector PowerPoint and PDF, using the same layout commands"}
    (ROOT / "verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    build()
    render_and_check()
